"""The yearly player-prop workbook, <Desktop>/<year>_nfl_props.xlsx, kept apart from the spread workbook.

Three tabs, rebuilt every run after the picks already in the workbook are read back (by header text), so nothing typed
in is lost:
  This Week   the season scoreboard (overall and by side) and this run's picks, best first: overs to bet now at
              DraftKings' opening line and, when enabled, unders to bet near kickoff; each with the line and where it
              opened, the forecast, the model's chance, the backtest win rate and the worst price still worth taking.
  Prop Log    every pick logged this season, graded from nflverse results, with its closing line.  Type the price you
              got in Your Price (yellow; blank counts as -110) and anything else in Notes.
  Model Info  the pick rule, the backtest record behind Hist. Win %, and how to read the workbook.
A pick is logged the first time it qualifies and kept as it was (a later run adds new picks only), so the log shows the
line that was available when the pick was made.  The file is saved to a temporary copy and swapped in, so a failed
save (for example with the workbook open in Excel) never leaves a half-written file.
"""

import math
import os
import shutil
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import nfl_sheet as sheet

SHEET_THIS_WEEK, SHEET_LOG, SHEET_INFO = 'This Week', 'Prop Log', 'Model Info'
NAVY, WHITE, GRAY_TEXT = '1F4E78', 'FFFFFF', '595959'
MANUAL_HEADER, MANUAL_CELL, DETAIL_HEADER = 'FFD966', 'FFF8E1', '7F7F7F'
WIN_FILL, LOSS_FILL, PICK_FILL, CARD_FILL = 'C6EFCE', 'FFC7CE', 'BDD7EE', 'F2F2F2'
AMERICAN = '+0;-0;0'

# Prop Log columns: what you look at, your entries, the grading, then hidden detail (keys and raw numbers)
LOG_LAYOUT = (
    ('week', 'Week', 6), ('kickoff', 'Kickoff (ET)', 14), ('game', 'Game', 12), ('player', 'Player', 22),
    ('prop', 'Prop', 17), ('pick', 'Pick', 11), ('open_line', 'Opened At', 9), ('forecast', 'Forecast', 9),
    ('chance', 'Model Chance', 9), ('hist', 'Hist. Win %', 9), ('bet_at', 'Bet At (worst price)', 10),
    ('price', 'Your Price', 9), ('result', 'Result', 8), ('outcome', 'W/L', 6), ('units', 'Units', 7),
    ('close', 'Closing Line', 9), ('clv', 'Line vs Close', 9), ('notes', 'Notes', 24), ('why', 'Why', 90),
    ('team', 'Team', 6), ('line', 'Line', 7), ('logged_at', 'Logged', 17), ('season', 'Season', 7),
    ('game_id', 'game_id', 16), ('gsis_id', 'gsis_id', 12), ('market', 'market', 12), ('side', 'side', 7),
)
COL = {key: i for i, (key, _, _) in enumerate(LOG_LAYOUT, start=1)}
LETTER = {key: get_column_letter(i) for key, i in COL.items()}
FIRST_DETAIL = COL['team']
MANUAL_FIELDS = ('price', 'notes')
KEY_FIELDS = ('game_id', 'gsis_id', 'market', 'side')
READ_HEADERS = {header: key for key, header, _ in LOG_LAYOUT}

THIS_WEEK_LAYOUT = (
    ('kickoff', 'Kickoff (ET)', 14), ('game', 'Game', 12), ('player', 'Player', 22), ('prop', 'Prop', 17),
    ('pick', 'Pick', 11), ('open_line', 'Opened At', 9), ('forecast', 'Forecast', 9), ('chance', 'Model Chance', 9),
    ('hist', 'Hist. Win %', 9), ('bet_at', 'Bet At (worst price)', 10), ('why', 'Why', 90),
)


def _fill(color):
    return PatternFill(fill_type='solid', start_color=color, end_color=color)


def _value(value):
    """Plain Python value for a cell: NaN and NaT become blank, numpy numbers become int/float."""
    if value is None:
        return None
    if hasattr(value, 'item') and not isinstance(value, str):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _header(cell, color=NAVY, font_color=WHITE):
    cell.font = Font(bold=True, color=font_color)
    cell.fill = _fill(color)
    cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')


def workbook_path(year, output_dir=None):
    output_dir = output_dir or sheet.get_desktop_path()
    os.makedirs(output_dir, exist_ok=True)
    return os.path.join(output_dir, f'{year}_nfl_props.xlsx')


def pick_key(record):
    return tuple(str(record.get(field)) for field in KEY_FIELDS)


def read_log(path):
    """The picks already in the workbook's Prop Log, as dicts keyed by LOG_LAYOUT field (empty if there is none)."""
    if not os.path.exists(path):
        return []
    try:
        wb = load_workbook(path)
    except Exception:
        backup = path.replace('.xlsx', f'_unreadable_{datetime.now():%Y%m%d_%H%M%S}.xlsx')
        shutil.copy2(path, backup)
        raise RuntimeError(f"{path} could not be read; copied it to {backup}")
    if SHEET_LOG not in wb.sheetnames:
        return []
    ws = wb[SHEET_LOG]
    fields = {READ_HEADERS[ws.cell(1, c).value]: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value in READ_HEADERS}
    records = []
    for r in range(2, ws.max_row + 1):
        record = {field: ws.cell(r, c).value for field, c in fields.items()}
        if record.get('game_id') and record.get('gsis_id'):
            records.append(record)
    return records


def _write_log(ws, records):
    for c, (key, header, width) in enumerate(LOG_LAYOUT, start=1):
        cell = ws.cell(1, c, header)
        _header(cell, MANUAL_HEADER if key in MANUAL_FIELDS else DETAIL_HEADER if c >= FIRST_DETAIL else NAVY,
                '000000' if key in MANUAL_FIELDS else WHITE)
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, record in enumerate(records, start=2):
        for key, _, _ in LOG_LAYOUT:
            if key == 'units':
                continue
            ws.cell(r, COL[key], _value(record.get(key)))
        outcome, price = f"{LETTER['outcome']}{r}", f"{LETTER['price']}{r}"
        ws.cell(r, COL['units'], f'=IF({outcome}="W",IF({price}="",100/110,IF({price}<0,-100/{price},{price}/100)),'
                                 f'IF({outcome}="L",-1,0))')
        for key in ('chance', 'hist'):
            ws.cell(r, COL[key]).number_format = '0.0%'
        for key in ('bet_at', 'price'):
            ws.cell(r, COL[key]).number_format = AMERICAN
        ws.cell(r, COL['units']).number_format = '+0.00;-0.00;0'
        ws.cell(r, COL['clv']).number_format = '+0.0;-0.0;0'
        for key in MANUAL_FIELDS:
            ws.cell(r, COL[key]).fill = _fill(MANUAL_CELL)
        if record.get('outcome') in ('W', 'L'):
            ws.cell(r, COL['outcome']).fill = _fill(WIN_FILL if record['outcome'] == 'W' else LOSS_FILL)
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LOG_LAYOUT))}{max(2, len(records) + 1)}"
    ws.column_dimensions.group(get_column_letter(FIRST_DETAIL), get_column_letter(len(LOG_LAYOUT)), hidden=True)


def _write_picks(ws, row, title, picks):
    """A titled table of picks starting at `row`; returns the next free row."""
    ws.cell(row, 1, title).font = Font(bold=True, size=12, color=NAVY)
    row += 1
    for c, (key, header, width) in enumerate(THIS_WEEK_LAYOUT, start=1):
        _header(ws.cell(row, c, header))
        ws.column_dimensions[get_column_letter(c)].width = width
    for pick in picks:
        row += 1
        for c, (key, _, _) in enumerate(THIS_WEEK_LAYOUT, start=1):
            cell = ws.cell(row, c, _value(pick.get(key)))
            if key in ('chance', 'hist'):
                cell.number_format = '0.0%'
            elif key == 'bet_at':
                cell.number_format = AMERICAN
                cell.font = Font(bold=True)
            if key == 'pick':
                cell.fill = _fill(PICK_FILL)
                cell.font = Font(bold=True)
    if not picks:
        row += 1
        ws.cell(row, 1, "None: no line met the rule (or no DraftKings lines are posted yet for the games left).")
    return row + 2


def _write_this_week(ws, year, week, picks, skipped, updated_at, note, sides):
    log = f"'{SHEET_LOG}'!"
    outcome, side = f"{log}{LETTER['outcome']}:{LETTER['outcome']}", f"{log}{LETTER['side']}:{LETTER['side']}"
    ws['A1'] = f"NFL player prop picks: week {week}, {year}"
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = (f"Updated {updated_at:%a %b %d %Y, %I:%M %p}. DraftKings lines via ESPN; the free feed has no prices, so "
                f"bet a pick only at the price in Bet At or better (check it in the app).")
    ws['A2'].font = Font(italic=True, color=GRAY_TEXT)
    last = week - 1 if isinstance(week, int) else 0
    labels = (('Season W-L-P', f'=COUNTIF({outcome},"W")&"-"&COUNTIF({outcome},"L")&"-"&COUNTIF({outcome},"P")'),
              ('Win %', f'=IFERROR(COUNTIF({outcome},"W")/(COUNTIF({outcome},"W")+COUNTIF({outcome},"L")),"")'),
              ('Units (your price, else -110)', f'=SUM({log}{LETTER["units"]}:{LETTER["units"]})'),
              ('Overs W-L', f'=COUNTIFS({side},"Over",{outcome},"W")&"-"&COUNTIFS({side},"Over",{outcome},"L")'),
              ('Unders W-L', f'=COUNTIFS({side},"Under",{outcome},"W")&"-"&COUNTIFS({side},"Under",{outcome},"L")'),
              ('Avg line vs close', f'=IFERROR(AVERAGE({log}{LETTER["clv"]}:{LETTER["clv"]}),"")'),
              (f'Week {last} W-L' if last else 'Last week W-L',
               f'=COUNTIFS({log}{LETTER["week"]}:{LETTER["week"]},{last},{outcome},"W")&"-"&'
               f'COUNTIFS({log}{LETTER["week"]}:{LETTER["week"]},{last},{outcome},"L")'))
    for i, (label, formula) in enumerate(labels):
        col = 1 + i
        ws.cell(4, col, label).font = Font(bold=True, color=GRAY_TEXT)
        ws.cell(4, col).alignment = Alignment(wrap_text=True)
        cell = ws.cell(5, col, formula)
        cell.font = Font(bold=True, size=12)
        cell.fill = _fill(CARD_FILL)
        if label == 'Win %':
            cell.number_format = '0.0%'
        elif label.startswith('Units') or label.startswith('Avg'):
            cell.number_format = '+0.00;-0.00;0'
    row = 7
    overs = [p for p in picks if p.get('side') == 'Over']
    unders = [p for p in picks if p.get('side') == 'Under']
    if 'Over' in sides:
        row = _write_picks(ws, row, f"Over picks ({len(overs)}): bet now, at DraftKings' opening line; once the line moves "
                                    f"the edge is gone", overs)
    if 'Under' in sides:
        row = _write_picks(ws, row, f"Under picks ({len(unders)}): best bet close to kickoff", unders)
    notes = []
    if skipped:
        notes.append(f"Not evaluated (already kicked off, or no ESPN event): {', '.join(skipped)}.")
    if note:
        notes.append(note)
    notes += [
        "Pick: the side to bet at the line shown.  Opened At: DraftKings' first line.  Forecast: the model's average "
        "result; Model Chance: how often similar forecasts went that way in past seasons.",
        "Hist. Win %: how often the rule's picks in this market won (overs at opening lines, unders at closing lines, 2024 "
        "on, shrunk toward all markets).  Bet At: the worst price that still returns 3% at that rate.  More on Model Info.",
    ]
    for text in notes:
        ws.cell(row, 1, text).font = Font(italic=True, color=GRAY_TEXT)
        row += 1
    ws.freeze_panes = 'A6'


def _write_info(ws, rule_lines, records, evidence):
    """`records`: [(title, {market: (rate, wins, losses)}, overall rate)], one table per rule."""
    ws['A1'] = 'How the prop picks are made'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    row = 3
    for text in rule_lines:
        ws.cell(row, 1, text)
        row += 1
    for title, rates, overall in records:
        row += 1
        ws.cell(row, 1, title).font = Font(bold=True, color=NAVY)
        row += 1
        for c, header in enumerate(('Market', 'Wins', 'Losses', 'Win %', 'Hist. Win % used'), start=1):
            _header(ws.cell(row, c, header))
        for market, (rate, wins, losses) in rates.items():
            row += 1
            ws.cell(row, 1, market)
            ws.cell(row, 2, wins)
            ws.cell(row, 3, losses)
            ws.cell(row, 4, wins / (wins + losses) if wins + losses else None).number_format = '0.0%'
            ws.cell(row, 5, _value(rate)).number_format = '0.0%'
        row += 1
        ws.cell(row, 1, 'All markets').font = Font(bold=True)
        ws.cell(row, 4, _value(overall)).number_format = '0.0%'
        row += 1
    row += 1
    ws.cell(row, 1, "Evidence (python nfl_research_props.py)").font = Font(bold=True, color=NAVY)
    for text in evidence:
        row += 1
        ws.cell(row, 1, text)
    ws.column_dimensions['A'].width = 22
    for col in 'BCDE':
        ws.column_dimensions[col].width = 14


def write_workbook(path, year, week, picks, records, rule_records, rule_lines, evidence, skipped=(), note=None,
                   sides=('Over', 'Under')):
    """Write the three tabs to `path` (see the module docstring).  `records` is the whole Prop Log, in order;
    `rule_records` the backtest tables for Model Info; `sides` which pick sections This Week shows."""
    wb = Workbook()
    this_week = wb.active
    this_week.title = SHEET_THIS_WEEK
    _write_this_week(this_week, year, week, picks, list(skipped), datetime.now(), note, sides)
    _write_log(wb.create_sheet(SHEET_LOG), records)
    _write_info(wb.create_sheet(SHEET_INFO), rule_lines, rule_records, evidence)
    wb.active = 0
    temporary = path + '.saving'
    try:
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
    return path
